"""Offline smoke test: exercises config, fetch (field mapping), metrics, and the
full build.run() orchestration with every network call monkeypatched out. No real
HTTP calls happen anywhere in this file."""
import sys
from unittest.mock import patch

import pandas as pd

sys.path.insert(0, str(__import__('pathlib').Path(__file__).resolve().parent.parent))
from scanner import config, fetch, metrics, build  # noqa: E402

print("SEASON_START..SEASON_END:", config.SEASON_START, config.SEASON_END)
print("LEVELS:", config.LEVELS)
assert config.is_final_season(2023) is True
assert config.is_final_season(config.CURRENT_YEAR) is False
print("is_final_season() OK")

# --- determine_seasons_to_fetch ------------------------------------------------------
all_seasons = set(range(config.SEASON_START, config.SEASON_END + 1))
final_seasons = {s for s in all_seasons if config.is_final_season(s)}

# Cold start: nothing marked complete -> fetch everything.
cold = build.determine_seasons_to_fetch(completed_seasons=set())
assert cold == sorted(all_seasons), cold
print("Cold start fetch list OK:", cold)

# Warm start: all final seasons already complete -> fetch only current season.
warm = build.determine_seasons_to_fetch(completed_seasons=final_seasons)
assert warm == [config.SEASON_END], warm
print("Warm start fetch list OK (only current season):", warm)

# --- field mapping + row builder (reuse fixtures) -------------------------------------
fake_batting_line = {
    "teamAbbrev": "TUL", "teamName": "Tulsa Drillers", "leagueId": 109,
    "leagueName": "Texas League", "playerId": 663734, "playerFullName": "Test Hitter",
    "playerFirstName": "Test", "playerLastName": "Hitter", "playerUseName": "Test",
    "playerInitLastName": "T Hitter", "primaryPositionAbbrev": "C", "positionAbbrev": "C",
    "gamesPlayed": 90, "plateAppearances": 380, "atBats": 330, "hits": 95,
    "doubles": 20, "triples": 2, "homeRuns": 14, "rbi": 55, "runs": 50,
    "stolenBases": 3, "caughtStealing": 1, "baseOnBalls": 40, "intentionalWalks": 1,
    "strikeOuts": 80, "totalBases": 161, "gidpOpp": 20, "groundIntoDoublePlay": 8,
    "sacBunts": 1, "sacFlies": 3, "hitByPitch": 5, "extraBaseHits": 36,
    "groundOuts": 100, "airOuts": 90, "flyOuts": 60, "popOuts": 15, "lineOuts": 15,
    "catchersInterference": 0, "leftOnBase": 120, "groundHits": 40, "flyHits": 30,
    "popHits": 5, "lineHits": 20, "numberOfPitches": 1500, "totalSwings": 700,
    "swingAndMisses": 120, "ballsInPlay": 500, "reachedOnError": 4, "walkOffs": 1,
}
row = fetch._build_row(fake_batting_line, fetch.BATTING_FIELD_MAP, config.SEASON_END, 999, "AA")
bdf = metrics.add_batting_rates(pd.DataFrame([row]))
assert bdf.loc[0, "batting_AVG"] == round(95 / 330, 3)
print("Field map + metrics OK. batting_AVG =", bdf.loc[0, "batting_AVG"])

# --- full build.run() with network mocked out ------------------------------------------
def fake_pull_season_level(season, level, stats_type):
    """Return one synthetic row per call so we can trace the whole pipeline."""
    if stats_type == "batting":
        line = dict(fake_batting_line)
        line["playerId"] = 111
        line["playerFullName"] = "Active Catcher"
        return pd.DataFrame([fetch._build_row(line, fetch.BATTING_FIELD_MAP, season, 1, level)])
    else:
        line = {
            "teamAbbrev": "TUL", "teamName": "Tulsa Drillers", "leagueId": 109,
            "leagueName": "Texas League", "playerId": 222, "playerFullName": "Active Pitcher",
            "playerFirstName": "Active", "playerLastName": "Pitcher", "playerUseName": "Active",
            "playerInitLastName": "A Pitcher", "primaryPositionAbbrev": "P", "positionAbbrev": "P",
            "wins": 5, "losses": 3, "gamesPitched": 15, "gamesStarted": 15, "gamesFinished": 0,
            "completeGames": 0, "qualityStarts": 8, "shutouts": 0, "saveOpportunities": 0,
            "saves": 0, "holds": 0, "blownSaves": 0, "inningsPitched": "80.0", "outs": 240,
            "battersFaced": 340, "atBats": 300, "runs": 35, "hits": 70, "doubles": 12,
            "triples": 1, "earnedRuns": 30, "homeRuns": 8, "totalBases": 110, "baseOnBalls": 25,
            "intentionalWalks": 0, "strikeOuts": 90, "hitByPitch": 3, "balks": 0,
            "groundIntoDoublePlay": 5, "gidpOpp": 8, "catchersInterference": 0,
            "inheritedRunners": 0, "inheritedRunnersScored": 0, "bequeathedRunners": 1,
            "bequeathedRunnersScored": 0, "runSupport": 40, "sacFlies": 2, "stolenBases": 3,
            "caughtStealing": 1, "pickoffs": 0, "flyHits": 20, "popHits": 3, "lineHits": 15,
            "flyOuts": 50, "groundOuts": 80, "airOuts": 60, "popOuts": 10, "lineOuts": 8,
            "numberOfPitches": 1200, "totalSwings": 600, "swingAndMisses": 140,
            "ballsInPlay": 400, "strikes": 750, "wildPitches": 2,
        }
        return pd.DataFrame([fetch._build_row(line, fetch.PITCHING_FIELD_MAP, season, 1, level)])


pushed = {}
pushed_dfs = {}
def fake_push_rows(sheet_name, df):
    pushed[sheet_name] = len(df)
    pushed_dfs[sheet_name] = df
    return True

marked_complete = []
def fake_mark_complete(season):
    marked_complete.append(season)
    return True

with patch("scanner.build.fetch.pull_season_level", side_effect=fake_pull_season_level), \
     patch("scanner.build.fetch.fetch_player_bios", return_value=pd.DataFrame(
         [{"player_id": 111, "birth_date": "2001-04-12", "height": "6' 0\"", "weight": 180,
           "bats": "R", "throws": "R", "mlb_debut_date": None, "birth_country": "USA"},
          {"player_id": 222, "birth_date": "1999-09-30", "height": "6' 3\"", "weight": 210,
           "bats": "R", "throws": "R", "mlb_debut_date": None, "birth_country": "USA"}])), \
     patch("scanner.build.sheets_sync.get_completed_seasons", return_value=set()), \
     patch("scanner.build.sheets_sync.push_rows", side_effect=fake_push_rows), \
     patch("scanner.build.sheets_sync.mark_season_complete", side_effect=fake_mark_complete):
    build.run()

print("Pushed row counts per sheet:", pushed)
print("Seasons marked complete:", sorted(marked_complete))

n_seasons = len(all_seasons)
n_levels = len(config.LEVELS)
assert pushed["Catcher"] == n_seasons * n_levels, pushed  # our fake hitter is always a catcher
assert pushed["Batter"] == 0, pushed
assert pushed["Pitcher"] == n_seasons * n_levels, pushed
assert sorted(marked_complete) == sorted(final_seasons), marked_complete

assert "age_as_of_season" in pushed_dfs["Catcher"].columns
assert "age_as_of_season" in pushed_dfs["Pitcher"].columns
assert pushed_dfs["Catcher"]["birth_country"].iloc[0] == "USA"
print("Age/bio columns confirmed present in the actual pushed rows (not just the merge test).")

import openpyxl
wb = openpyxl.load_workbook(config.OUTPUT_DIR / "milb_prospect_scan_delta.xlsx")
assert wb.sheetnames == ["Batter", "Pitcher", "Catcher"], wb.sheetnames
print("Local artifact OK, sheets:", wb.sheetnames)

print("\nALL OFFLINE END-TO-END CHECKS PASSED")

# --- age/bio enrichment: prove it's ID-keyed, not name-keyed --------------------------
# Two players who happen to share a name, on purpose — this is the exact failure mode
# ("mapping to another person") the merge must NOT fall into.
import pandas as pd  # noqa: E402  (re-import fine, already imported above)

same_name_rows = pd.DataFrame([
    {"player_id": 501, "player_full_name": "Jose Rodriguez", "season": 2026,
     "batting_AB": 300, "batting_H": 90, "batting_PA": 340, "batting_BB": 30,
     "batting_HBP": 2, "batting_SF": 3, "batting_TB": 150, "batting_SO": 60,
     "batting_HR": 10, "batting_AO": 80, "batting_GO": 90},
    {"player_id": 999, "player_full_name": "Jose Rodriguez", "season": 2026,
     "batting_AB": 280, "batting_H": 70, "batting_PA": 310, "batting_BB": 20,
     "batting_HBP": 1, "batting_SF": 2, "batting_TB": 100, "batting_SO": 70,
     "batting_HR": 5, "batting_AO": 70, "batting_GO": 85},
])

fake_bios = pd.DataFrame([
    {"player_id": 501, "birth_date": "2000-05-01", "height": "6' 1\"", "weight": 190,
     "bats": "R", "throws": "R", "mlb_debut_date": None, "birth_country": "Dominican Republic"},
    {"player_id": 999, "birth_date": "2003-11-20", "height": "5' 11\"", "weight": 175,
     "bats": "L", "throws": "R", "mlb_debut_date": None, "birth_country": "Venezuela"},
])

aged = metrics.add_age(same_name_rows, fake_bios)
row_501 = aged[aged["player_id"] == 501].iloc[0]
row_999 = aged[aged["player_id"] == 999].iloc[0]

assert row_501["birth_date"] == "2000-05-01" and row_501["birth_country"] == "Dominican Republic"
assert row_999["birth_date"] == "2003-11-20" and row_999["birth_country"] == "Venezuela"
assert row_501["age_as_of_season"] > row_999["age_as_of_season"], (
    "Older player_id 501 should show as older than 999 -- if this fails, "
    "identity got crossed."
)
print(
    "Same-name collision test OK: two 'Jose Rodriguez' rows kept fully separate "
    "birth dates/countries/ages via player_id, never confused by name ->",
    f"501: age {row_501['age_as_of_season']}, {row_501['birth_country']} | "
    f"999: age {row_999['age_as_of_season']}, {row_999['birth_country']}"
)

# --- enrich_with_age() end-to-end, with the network call mocked out ------------------
with patch("scanner.build.fetch.fetch_player_bios", return_value=fake_bios):
    b_out, p_out, c_out = build.enrich_with_age(same_name_rows, pd.DataFrame(), pd.DataFrame())
assert "age_as_of_season" in b_out.columns
assert set(b_out["birth_country"]) == {"Dominican Republic", "Venezuela"}
print("build.enrich_with_age() OK -- age/bio columns attached correctly per player_id")

print("\nALL AGE/IDENTITY CHECKS PASSED")
